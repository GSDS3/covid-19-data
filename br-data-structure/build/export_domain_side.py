# -*- coding: utf-8 -*-
"""
데이터 트리 + 관측 카탈로그 → Domain_Side 양식 엑셀 (dist/Domain_Side_사업보고서.xlsx)

시트
  사업보고서        Depth / Level 0~4 / Item / Observations  (사용자 양식. Level 4 = 항목 세분화)
  Observations      관측 한 건 = 한 행. 관측 규약 11필드 중 항목 단위로 확정되는 칸을 채운 것
  표 유형(근거)      구조 조사 5건에서 실제로 관측된 표 머리글. 관측·세분화의 근거
  정보트리 매핑      데이터 노드 → 정보 트리 노드 (주제 연결/항목 대응)
  Depth 정렬(5건)    구조 조사 5건의 절 트리를 Depth 별로 나래비 세운 결과(공통/일부/단독/미관찰)
  범례              열 정의·어휘·세분화 규칙

필요: openpyxl  (pip install openpyxl)
사용: python3 build/export_domain_side.py [--out 경로.xlsx]
"""
import os, json, argparse, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
HDR = PatternFill("solid", fgColor="1F3A5F"); HDRF = Font(color="FFFFFF", bold=True, size=10)
L1F = PatternFill("solid", fgColor="DCE6F1"); L2F = PatternFill("solid", fgColor="EFF4FA")
SUBF = PatternFill("solid", fgColor="FFF6E5")
THIN = Border(*[Side(style="thin", color="D9D9D9")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

def load(p): return json.load(open(os.path.join(ROOT, p), encoding="utf-8"))

def style_header(ws, row=1, freeze=None, widths=(), wrapcols=()):
    for c in ws[row]:
        if c.value is None: continue
        c.fill = HDR; c.font = HDRF; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        if w: ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze or ws.cell(row=row + 1, column=1).coordinate
    ws.auto_filter.ref = f"A{row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    for r in ws.iter_rows(min_row=row + 1):
        for c in r:
            c.alignment = WRAP if c.column in wrapcols else TOP
            c.border = THIN
            if c.font is None or not c.font.size: c.font = Font(size=10)

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default=os.path.join(ROOT, "dist", "Domain_Side_사업보고서.xlsx"))
    a = ap.parse_args()
    tree = load("skeleton/data_tree_template.json"); obs = load("skeleton/observations.json")
    mapping = load("skeleton/mapping.json"); info = load("skeleton/info_tree.json")
    skel = load("dist/skeleton.json") if os.path.exists(os.path.join(ROOT, "dist", "skeleton.json")) else {}
    items = obs["items"]

    order = []
    def walk(n):
        order.append(n)
        for c in n.get("children") or []: walk(c)
    walk(tree["root"])

    wb = Workbook(); ws = wb.active; ws.title = "사업보고서"
    ws.append(["", "사업보고서 Data Tree", "", "", "", "", "", "", "", "", "", ""])
    ws["B1"].font = Font(bold=True, size=13)
    ws.append(["", "Depth", "Level 0", "Level 1", "Level 2", "Level 3", "Level 4", "Item", "Observations",
               "관측 수", "적용 범위", "근거(구조조사 5건)"])
    scope_ko = {"common": "공통", "conditional": "조건부"}
    for n in order:
        nid = n["id"]; parts = nid.split(".")           # DT, A001, L1, L2, L3
        lv = ["A001"] + parts[2:5] + [""] * 5
        depth = len(parts) - 2
        it = items.get(nid)
        sc = n.get("scope") or ""; sc = scope_ko.get(sc, sc.replace("industry:", "업종:"))
        names = []
        if it and it["subitem_count"] == 0:
            names = [o["name"] for g in it["groups"] for o in g["observations"]]
        ev = ""
        if it: ev = ("/".join(x[-2:] for x in it["evidence_reports"]) + (" (절)" if it["evidence_level"] == "절" else "")) if it["evidence_reports"] else ""
        ws.append(["", depth, lv[0], lv[1], lv[2], lv[3], "", n["title"],
                   " · ".join(names) if names else ("아래 하위 항목으로 세분" if it and it["subitem_count"] else ""),
                   it["observation_count"] if it else "", sc, ev])
        r = ws.max_row
        if depth == 1:
            for c in ws[r]: c.fill = L1F
            ws.cell(row=r, column=8).font = Font(bold=True, size=10)
        elif depth == 2 and (n.get("children") or []):
            for c in ws[r]: c.fill = L2F
        if it and it["subitem_count"]:
            for g in it["groups"]:
                if not g["title"]: continue
                ws.append(["", depth + 1, lv[0], lv[1], lv[2], lv[3], g["code"], g["title"],
                           " · ".join(o["name"] for o in g["observations"]), len(g["observations"]), sc, ev])
                for c in ws[ws.max_row]: c.fill = SUBF
    style_header(ws, row=2, freeze="C3", widths=(3, 7, 9, 9, 9, 9, 9, 46, 74, 8, 11, 17), wrapcols=(8, 9, 12))
    ws.auto_filter.ref = f"B2:L{ws.max_row}"

    # ── Observations
    ws2 = wb.create_sheet("Observations")
    ws2.append(["관측 ID", "노드 ID", "Depth", "Level 1", "Level 2", "Level 3", "Level 4", "Item", "하위 항목",
                "관측명(측정량)", "측정유형", "단위", "관측 축(행 축·주체)", "시간 역할", "범위·측정기준", "조건·주의", "근거"])
    for n in order:
        it = items.get(n["id"])
        if not it: continue
        parts = n["id"].split("."); lv = parts[2:5] + [""] * 3
        ev = ("/".join(x[-2:] for x in it["evidence_reports"]) + (" (절)" if it["evidence_level"] == "절" else "")) if it["evidence_reports"] else "서식"
        for g in it["groups"]:
            for o in g["observations"]:
                ws2.append([o["id"], n["id"], len(parts) - 2 + (1 if g["title"] else 0), lv[0], lv[1], lv[2],
                            g["code"] or "", n["title"], g["title"] or "", o["name"], o["type"], o["unit"],
                            o["axis"], o["time"], o["basis"], o["note"], ev])
    style_header(ws2, widths=(17, 21, 6, 8, 8, 8, 8, 34, 26, 34, 10, 13, 30, 17, 30, 46, 13), wrapcols=(8, 9, 10, 13, 14, 15, 16))

    # ── 표 유형(근거)
    ws3 = wb.create_sheet("표 유형(근거)")
    ws3.append(["노드 ID", "절·항목", "표 캡션(원문)", "머리글 시그니처", "단위 표기", "보고서 수", "보고서", "발생 수", "병합 비율"])
    tnode = {n["id"]: n["title"] for n in order}
    for t in sorted(skel.get("table_types") or [], key=lambda x: (x["node"], -len(x["reports"]), -x["count"])):
        ws3.append([t["node"], tnode.get(t["node"], ""), t.get("caption") or "", t.get("header") or "", t.get("unit") or "",
                    len(t["reports"]), "/".join(x[-2:] for x in t["reports"]), t["count"], round(t.get("merged_ratio") or 0, 2)])
    style_header(ws3, widths=(21, 30, 56, 60, 15, 9, 13, 8, 9), wrapcols=(3, 4))

    # ── 정보트리 매핑
    ws4 = wb.create_sheet("정보트리 매핑")
    ws4.append(["데이터 노드 ID", "Item", "정보트리 노드", "정보트리 항목", "대응 수준", "요구 스킬"])
    iflat = {}
    def iwalk(n):
        iflat[n["id"]] = n
        for c in n.get("children") or []: iwalk(c)
    for r in (info["roots"] if "roots" in info else [info["root"]]): iwalk(r)
    for m in mapping["edges"] if "edges" in mapping else mapping.get("mappings", []):
        d = m.get("data"); tgt = m.get("info") or m.get("its")
        for t in (tgt if isinstance(tgt, list) else [tgt]):
            node = iflat.get(t, {})
            ws4.append([d, tnode.get(d, ""), t, node.get("title", ""), m.get("level", ""), " · ".join(node.get("skills") or [])])
    style_header(ws4, widths=(21, 38, 14, 38, 11, 30), wrapcols=(2, 4, 6))

    # ── Depth 정렬(구조 조사 5건)
    ap_path = os.path.join(ROOT, "inputs", "derived", "depth_alignment.json")
    if os.path.exists(ap_path):
        al = json.load(open(ap_path, encoding="utf-8"))
        rid = [r["id"] for r in al["reports"]]
        ws6 = wb.create_sheet("Depth 정렬(5건)")
        ws6.append(["Depth", "노드 ID", "Item", "분류", "커버리지", "제목 표기 변형"] +
                   [f'{r["id"]} {r["company"]} {r["fiscal_year"]}' for r in al["reports"]] + ["표 합계", "셀 합계"])
        for r in al["rows"]:
            cells, tables = 0, 0
            per = []
            for b in rid:
                v = r["per_report"].get(b)
                if v:
                    tables += v["tables"]; cells += v["cells"]
                    per.append(f'{v["title"]}' + (f' [표{v["tables"]}]' if v["tables"] else ""))
                else:
                    per.append("")
            ws6.append([r["level"], r["node"], r["title"], r["class"], r["coverage"],
                        len(r["title_variants"])] + per + [tables, cells])
            row = ws6.max_row
            if r["class"] == "공통":
                ws6.cell(row=row, column=4).fill = PatternFill("solid", fgColor="E2F0D9")
            elif r["class"] == "미관찰":
                ws6.cell(row=row, column=4).fill = PatternFill("solid", fgColor="F2F2F2")
            else:
                ws6.cell(row=row, column=4).fill = PatternFill("solid", fgColor="FFF2CC")
            if r["level"] == 1:
                for c in ws6[row]: c.font = Font(bold=True, size=10)
        style_header(ws6, widths=(7, 21, 34, 9, 10, 11) + (30,) * len(rid) + (9, 10),
                     wrapcols=tuple(range(7, 7 + len(rid))))
        s6 = al["summary"]
        ws6.append([])
        ws6.append(["", "요약", " / ".join(f'L{l}: 총 {s6[f"L{l}"]["총"]} · 전건공통 {s6[f"L{l}"]["공통"]} · 일부 {s6[f"L{l}"]["일부"]} · 단독 {s6[f"L{l}"]["단독"]} · 미관찰 {s6[f"L{l}"]["미관찰"]}' for l in (1, 2, 3))])
        ws6.cell(row=ws6.max_row, column=2).font = Font(bold=True)

    # ── 범례
    ws5 = wb.create_sheet("범례")
    for row in [
        ["사업보고서 공통 데이터체계 — 양식 설명"], [],
        ["1. 층위"],
        ["Level 0", "공시서식 코드 A001(사업보고서). 다른 정기공시(A002 반기·A003 분기)는 같은 방식으로 별도 트리를 만든다."],
        ["Level 1", "장(章). COVER·CONF·I~XII·EXP 15개. 구조 조사 5건 모두에서 동일했다."],
        ["Level 2", "절(節). 장에 따라 절이 없으면(IV·VII·IX·X·EXP) 항목이 이 자리에 온다."],
        ["Level 3", "항목. 서식이 정한 기재 단위."],
        ["Level 4", "항목 세분화. 한 항목 안에 grain(행 축)이 다른 표가 둘 이상일 때만 만든다."], [],
        ["2. 세분화와 관측의 경계"],
        ["세분화(Level 4)", "행 축이 다른 표 → 별개 하위 항목. 예: 신용평가 실적 표와 신용등급 정의 표."],
        ["관측(Observation)", "같은 표 안의 서로 다른 측정량 → 관측 여러 건. 예: 생산실적과 가동률."], [],
        ["3. Observations 열 정의"],
        ["관측명(측정량)", "무엇을 잰 값인가. 원문의 열 이름(라벨)과 구별한다."],
        ["측정유형", " / ".join(sorted(obs["vocab"]["measure_types"]))],
        ["단위", "통화·배율(원/천원/백만원/천USD)·수량(주·톤)·비율(%)·배수(회). 표마다 다르므로 값에 붙여 보존한다."],
        ["관측 축(행 축·주체)", "그 값을 한 건으로 특정하는 키. 행 축과 주체 역할(보고회사·종속·상대방)을 적는다."],
        ["시간 역할", "시점(기말)·기간(사업연도)·사건일(결의·승인·효력)·공시일을 구분한다."],
        ["범위·측정기준", "연결/별도, 내부거래 포함 여부, 장부/공정가치, 산식의 분모 등."],
        ["조건·주의", "값의 의미를 뒤집을 수 있는 각주·예외, 그리고 5건 조사에서 실제로 발견한 함정."],
        ["근거", "그 항목의 표가 실제로 관측된 보고서. '(절)'은 절 단위까지만 확인된 것. '서식'은 기업공시서식 작성기준에 따른 초안."], [],
        ["4. 값의 상태 어휘"], ["", " / ".join(obs["vocab"]["status"])],
        ["", "실제와 계획·예정·잠정·가정을 같은 열에 섞지 않는다. 한도(약정·승인금액)는 잔액·지급액이 아니다."], [],
        ["5. 수치 규모"],
        ["항목", str(obs["meta"]["item_count"])], ["하위 항목(Level 4)", str(obs["meta"]["subitem_count"])],
        ["관측", str(obs["meta"]["observation_count"])],
        ["표 유형 근거", f"{len(skel.get('table_types') or [])}종 (구조 조사 5건)"], [],
        ["6. 확정과 초안"],
        ["확정", "Level 1·2 구조. 구조 조사 5건의 목차에서 동일하게 확인했다(보험업 II장만 변형)."],
        ["초안", "Level 3·4 항목과 관측 정의. 서식 기준과 5건 관측을 합친 것이며, 나머지 보고서 원문으로 검증·개정한다."],
        ["미검증", "업종 변형(은행·증권·건설·지주·공기업)과 XBRL 표준계정 대응."],
    ]: ws5.append(row)
    ws5.column_dimensions["A"].width = 24; ws5.column_dimensions["B"].width = 118
    for r in ws5.iter_rows():
        for c in r:
            c.alignment = WRAP
            if c.column == 1 and c.value and str(c.value)[:1].isdigit() and "." in str(c.value)[:3]: c.font = Font(bold=True, size=11)
    ws5["A1"].font = Font(bold=True, size=14)

    os.makedirs(os.path.dirname(a.out), exist_ok=True); wb.save(a.out)
    print(f"→ {a.out}")
    print(f"   사업보고서 {ws.max_row - 2}행 · Observations {ws2.max_row - 1}행 · 표 유형 {ws3.max_row - 1}행 · 매핑 {ws4.max_row - 1}행")

if __name__ == "__main__":
    main()
